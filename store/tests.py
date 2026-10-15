from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import RequestFactory, TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .context_processors import announcements as announcements_context
from .models import (
    Announcement, CreditEntry, Customer, Invoice, InvoiceItemBatch, Product, Shop,
    StockAdjustment, StockBatch, UserProfile,
)
from .services import calculate_credit_balance, cancel_invoice, create_invoice


class ShopTestCase(TestCase):
    def setUp(self):
        self.shop = Shop.objects.create(name="One Shop", setup_completed=True, invoice_prefix="ONE")
        self.other_shop = Shop.objects.create(name="Other Shop", setup_completed=True, invoice_prefix="TWO")
        self.user = User.objects.create_user("owner", password="Test-password-123")
        UserProfile.objects.create(user=self.user, shop=self.shop, role=UserProfile.Role.OWNER, must_change_password=False)
        self.customer = Customer.objects.create(shop=self.shop, name="Customer", whatsapp_number="9876543210")
        self.product = Product.objects.create(
            shop=self.shop, name="Tea", unit=Product.Unit.PCS, selling_price=Decimal("10.00"),
            gst_percent=0, low_stock_limit=Decimal("2"),
        )
        self.batch = StockBatch.objects.create(
            shop=self.shop, product=self.product, quantity_received=Decimal("10"),
            quantity_remaining=Decimal("10"),
        )


class TenantIsolationTests(ShopTestCase):
    def test_login_page_explains_shopsetu_purpose_in_hindi_and_english(self):
        response = self.client.get(reverse("login"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "आपकी दुकान के लिए आसान डिजिटल हिसाब")
        self.assertContains(response, "स्टॉक संभालने")
        self.assertContains(response, "A simpler way to manage your shop")
        self.assertContains(response, "manage stock")
        self.assertContains(response, "login-shop-illustration")
        self.assertContains(response, "prefers-reduced-motion")

    def test_dashboard_aggregates_expiry_and_credit_instead_of_per_customer_queries(self):
        second_customer = Customer.objects.create(
            shop=self.shop, name="Second Customer", whatsapp_number="9876543212",
        )
        CreditEntry.objects.create(
            shop=self.shop, customer=self.customer, entry_type=CreditEntry.EntryType.GIVEN,
            amount=Decimal("50.00"),
        )
        CreditEntry.objects.create(
            shop=self.shop, customer=self.customer, entry_type=CreditEntry.EntryType.RECEIVED,
            amount=Decimal("12.00"),
        )
        CreditEntry.objects.create(
            shop=self.shop, customer=second_customer, entry_type=CreditEntry.EntryType.GIVEN,
            amount=Decimal("10.00"),
        )
        self.batch.expiry_date = timezone.localdate() + timedelta(days=5)
        self.batch.save(update_fields=["expiry_date"])
        self.client.force_login(self.user)

        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(reverse("home"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["pending_credit"], Decimal("48.00"))
        self.assertEqual(response.context["expiry_counts"][7], 1)
        sql = [query["sql"].lower() for query in queries]
        self.assertEqual(sum("store_creditentry" in query for query in sql), 1)
        self.assertEqual(
            sum("store_stockbatch" in query and "count(" in query for query in sql), 1,
        )

    def test_announcement_schedule_and_shop_targets_are_respected(self):
        now = timezone.now()
        global_item = Announcement.objects.create(title="Global", message="All shops")
        targeted_item = Announcement.objects.create(title="Targeted", message="One shop")
        targeted_item.targeted_shops.add(self.other_shop)
        Announcement.objects.create(
            title="Future", message="Not yet", starts_at=now + timedelta(days=1),
        )
        Announcement.objects.create(
            title="Ended", message="Already ended", ends_at=now - timedelta(days=1),
        )
        request = RequestFactory().get("/")
        request.user = self.user

        visible = list(announcements_context(request)["active_announcements"])
        self.assertEqual(visible, [global_item])

        other_user = User.objects.create_user("other-owner")
        UserProfile.objects.create(
            user=other_user, shop=self.other_shop, role=UserProfile.Role.OWNER,
            must_change_password=False,
        )
        request.user = other_user
        self.assertEqual(
            set(announcements_context(request)["active_announcements"]),
            {global_item, targeted_item},
        )

    def test_inactive_shop_is_blocked_from_shopkeeper_pages(self):
        self.shop.is_active = False
        self.shop.save(update_fields=["is_active"])
        self.client.force_login(self.user)

        response = self.client.get(reverse("home"))
        self.assertRedirects(response, reverse("account_inactive"))

    def test_pwa_manifest_and_service_worker_are_public(self):
        manifest = self.client.get(reverse("manifest"))
        self.assertEqual(manifest.status_code, 200)
        self.assertEqual(manifest.json()["name"], "ShopSetu")
        self.assertEqual(self.client.get(reverse("service_worker")).status_code, 200)

    def test_product_and_stock_pages_exclude_other_shop(self):
        Product.objects.create(
            shop=self.other_shop, name="Secret item", unit=Product.Unit.PCS,
            selling_price=Decimal("99.00"),
        )
        self.client.force_login(self.user)
        self.assertRedirects(self.client.get(reverse("products")), reverse("stock"))
        response = self.client.get(reverse("stock"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Tea")
        self.assertContains(response, "Delete")
        self.assertNotContains(response, "Secret item")

    def test_product_search_is_shop_scoped_and_billing_search_submits_query(self):
        Product.objects.create(
            shop=self.other_shop, name="Tea biscuits", unit=Product.Unit.PCS,
            selling_price=Decimal("5.00"),
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("product_search"), {"q": "Tea"}, HTTP_HX_REQUEST="true")
        self.assertContains(response, "Tea")
        self.assertNotContains(response, "Tea biscuits")
        billing_response = self.client.get(reverse("billing"))
        self.assertContains(billing_response, 'name="q"')
        self.assertContains(billing_response, 'id="product-results"')
        self.assertContains(billing_response, "Type a product name above")
        self.assertNotContains(billing_response, "Tea")

    def test_credit_whatsapp_reminder_includes_the_customer_shop_name(self):
        CreditEntry.objects.create(
            shop=self.shop, customer=self.customer,
            entry_type=CreditEntry.EntryType.GIVEN, amount=Decimal("125.50"),
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("credit_whatsapp", args=[self.customer.pk]))

        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            response["Location"],
            f"https://wa.me/91{self.customer.whatsapp_number}"
            "?text=Hello+Customer%2C+your+pending+balance+at+One+Shop+is+Rs.+125.50.",
        )

    def test_shop_owned_manager_can_scope_to_authenticated_user(self):
        other = Product.objects.create(
            shop=self.other_shop, name="Other item", unit=Product.Unit.PCS,
            selling_price=Decimal("1.00"),
        )
        self.assertEqual(list(Product.objects.for_user(self.user)), [self.product])
        self.assertNotIn(other, Product.objects.for_user(self.user))

    def test_promotion_customer_remains_eligible_after_previous_message(self):
        self.customer.is_promotion_allowed = True
        self.customer.save(update_fields=["is_promotion_allowed"])
        self.client.force_login(self.user)

        response = self.client.post(reverse("promotion_send", args=[self.customer.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], f"https://wa.me/{self.customer.whatsapp_international}?text=Hello+Customer%2C+we+have+fresh+offers+for+you+at+our+shop%21")
        self.assertEqual(self.customer.promotion_logs.count(), 1)

        response = self.client.get(reverse("promotions"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.customer.name)
        self.assertContains(response, "no sending cooldown")
        self.assertContains(response, ">Send</button>")
        self.assertNotContains(response, "Share attachment")

        response = self.client.post(reverse("promotion_send", args=[self.customer.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.customer.promotion_logs.count(), 2)

    def test_promotion_send_supports_share_sheet_log_response_and_custom_message(self):
        self.customer.is_promotion_allowed = True
        self.customer.save(update_fields=["is_promotion_allowed"])
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("promotion_send", args=[self.customer.pk]),
            {"message": "Special offer for {name}!"},
            HTTP_ACCEPT="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "shared"})
        log = self.customer.promotion_logs.get()
        self.assertEqual(log.message, "Special offer for Customer!")

    def test_promotions_page_offers_image_and_pdf_attachment_without_cooldown(self):
        self.customer.is_promotion_allowed = True
        self.customer.save(update_fields=["is_promotion_allowed"])
        self.client.force_login(self.user)
        response = self.client.get(reverse("promotions"))
        self.assertContains(response, 'type="file"')
        self.assertContains(response, "image/*,application/pdf")
        self.assertContains(response, "attach the selected file there")
        self.assertContains(response, "no sending cooldown")

    def test_owner_pages_render(self):
        self.client.force_login(self.user)
        urls = [
            reverse("home"), reverse("setup"),
            reverse("product_edit", args=[self.product.pk]),
            reverse("stock"), reverse("stock_add"), reverse("stock_bulk_import"),
            reverse("stock_bulk_template"), reverse("stock_correct", args=[self.batch.pk]),
            reverse("billing"), reverse("customers"), reverse("customer_create"),
            reverse("customer_edit", args=[self.customer.pk]), reverse("customer_detail", args=[self.customer.pk]),
            reverse("credit_list"), reverse("credit_add", args=[self.customer.pk]),
            reverse("promotions"), reverse("staff"), reverse("staff_create"), reverse("reports"),
        ]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_legacy_product_add_route_redirects_to_stock_add(self):
        self.client.force_login(self.user)
        self.assertRedirects(self.client.get(reverse("product_create")), reverse("stock_add"))

    def test_owner_can_delete_product_and_its_batches(self):
        product = Product.objects.create(
            shop=self.shop, name="Unused", unit=Product.Unit.PCS,
            selling_price=Decimal("5.00"),
        )
        batch = StockBatch.objects.create(
            shop=self.shop, product=product, quantity_received=Decimal("4"),
            quantity_remaining=Decimal("4"),
        )
        self.client.force_login(self.user)

        response = self.client.post(reverse("product_delete", args=[product.pk]))

        self.assertRedirects(response, reverse("stock"))
        self.assertFalse(Product.objects.filter(pk=product.pk).exists())
        self.assertFalse(StockBatch.objects.filter(pk=batch.pk).exists())

    def test_product_delete_is_post_only_and_shop_scoped(self):
        other_product = Product.objects.create(
            shop=self.other_shop, name="Other product", unit=Product.Unit.PCS,
            selling_price=Decimal("5.00"),
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("product_delete", args=[self.product.pk]))
        self.assertEqual(response.status_code, 405)
        response = self.client.post(reverse("product_delete", args=[other_product.pk]))
        self.assertEqual(response.status_code, 404)
        self.assertTrue(Product.objects.filter(pk=self.product.pk).exists())
        self.assertTrue(Product.objects.filter(pk=other_product.pk).exists())

    def test_staff_cannot_delete_product(self):
        staff = User.objects.create_user("staff")
        UserProfile.objects.create(
            user=staff, shop=self.shop, role=UserProfile.Role.STAFF,
            must_change_password=False,
        )
        self.client.force_login(staff)

        response = self.client.post(reverse("product_delete", args=[self.product.pk]))

        self.assertEqual(response.status_code, 403)
        self.assertTrue(Product.objects.filter(pk=self.product.pk).exists())

    def test_product_with_invoice_history_cannot_be_deleted(self):
        create_invoice(
            user=self.user, shop=self.shop, customer=self.customer,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "1"}]' % self.product.pk,
        )
        self.client.force_login(self.user)

        response = self.client.post(reverse("product_delete", args=[self.product.pk]), follow=True)

        self.assertContains(response, "cannot be deleted because it has invoice")
        self.assertTrue(Product.objects.filter(pk=self.product.pk).exists())
        self.assertTrue(StockBatch.objects.filter(pk=self.batch.pk).exists())

    def test_product_with_adjustment_history_cannot_be_deleted(self):
        StockAdjustment.objects.create(
            shop=self.shop, batch=self.batch, quantity_change=Decimal("-1"),
            reason=StockAdjustment.Reason.DAMAGED, created_by=self.user,
        )
        self.client.force_login(self.user)

        response = self.client.post(reverse("product_delete", args=[self.product.pk]), follow=True)

        self.assertContains(response, "cannot be deleted because it has invoice or stock-adjustment history")
        self.assertTrue(Product.objects.filter(pk=self.product.pk).exists())
        self.assertTrue(StockBatch.objects.filter(pk=self.batch.pk).exists())

    def test_credit_payment_cannot_exceed_pending_balance(self):
        CreditEntry.objects.create(
            shop=self.shop, customer=self.customer,
            entry_type=CreditEntry.EntryType.GIVEN, amount=Decimal("100.00"),
        )
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("credit_add", args=[self.customer.pk]),
            {
                "entry_type": CreditEntry.EntryType.RECEIVED,
                "amount": "100.01",
                "note": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "cannot exceed the pending credit of Rs. 100.00")
        self.assertFalse(CreditEntry.objects.filter(
            shop=self.shop, customer=self.customer,
            entry_type=CreditEntry.EntryType.RECEIVED,
        ).exists())
        self.assertEqual(calculate_credit_balance(self.customer), Decimal("100.00"))

    def test_credit_payment_can_settle_exact_pending_balance(self):
        CreditEntry.objects.create(
            shop=self.shop, customer=self.customer,
            entry_type=CreditEntry.EntryType.GIVEN, amount=Decimal("100.00"),
        )
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("credit_add", args=[self.customer.pk]),
            {
                "entry_type": CreditEntry.EntryType.RECEIVED,
                "amount": "100.00",
                "note": "",
            },
        )

        self.assertRedirects(response, reverse("customer_detail", args=[self.customer.pk]))
        self.assertEqual(calculate_credit_balance(self.customer), Decimal("0.00"))

    def test_stock_add_creates_product_and_opening_batch_together(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse("stock_add"), {
            "action": "create_product",
            "name": "Coffee",
            "unit": Product.Unit.PACKET,
            "selling_price": "25.00",
            "purchase_price": "18.00",
            "gst_percent": "5",
            "hsn_code": "0901",
            "low_stock_limit": "3",
            "quantity_received": "12",
            "expiry_date": "",
            "note": "Opening inventory",
        })

        self.assertRedirects(response, reverse("stock"))
        product = Product.objects.get(shop=self.shop, name="Coffee")
        batch = StockBatch.objects.get(shop=self.shop, product=product)
        self.assertEqual(batch.quantity_received, Decimal("12"))
        self.assertEqual(batch.quantity_remaining, Decimal("12"))
        self.assertEqual(batch.purchase_price, Decimal("18"))
        self.assertEqual(batch.note, "Opening inventory")

    def test_stock_add_duplicate_product_does_not_create_product_or_batch(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse("stock_add"), {
            "action": "create_product",
            "name": "Tea",
            "unit": Product.Unit.PCS,
            "selling_price": "10.00",
            "purchase_price": "",
            "gst_percent": "0",
            "hsn_code": "",
            "low_stock_limit": "0",
            "quantity_received": "5",
            "expiry_date": "",
            "note": "",
        })

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "already exists")
        self.assertEqual(Product.objects.filter(shop=self.shop, name="Tea").count(), 1)
        self.assertEqual(StockBatch.objects.filter(shop=self.shop, product=self.product).count(), 1)


class BillingServiceTests(ShopTestCase):
    def test_billing_creates_new_customer_and_credit_invoice_together(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse("billing"), {
            "new_customer_name": "New buyer",
            "new_customer_whatsapp_number": "9876543211",
            "new_customer_address": "Main road",
            "payment_type": Invoice.PaymentType.CREDIT,
            "amount_paid": "0",
            "items_json": '[{"product_id": %s, "quantity": "2"}]' % self.product.pk,
        })

        customer = Customer.objects.get(shop=self.shop, whatsapp_number="9876543211")
        invoice = Invoice.objects.get(shop=self.shop)
        self.assertRedirects(response, reverse("invoice_detail", args=[invoice.pk]))
        self.assertEqual(invoice.customer, customer)
        self.assertEqual(calculate_credit_balance(customer), Decimal("20.00"))

    def test_failed_bill_rolls_back_inline_customer_creation(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse("billing"), {
            "new_customer_name": "New buyer",
            "new_customer_whatsapp_number": "9876543211",
            "payment_type": Invoice.PaymentType.CREDIT,
            "amount_paid": "0",
            "items_json": '[{"product_id": %s, "quantity": "11"}]' % self.product.pk,
        })

        self.assertEqual(response.status_code, 200)
        self.assertFalse(Customer.objects.filter(shop=self.shop, whatsapp_number="9876543211").exists())
        self.assertFalse(Invoice.objects.filter(shop=self.shop).exists())
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.quantity_remaining, Decimal("10"))

    def test_fefo_deducts_earliest_expiry_and_skips_expired_batch(self):
        today = timezone.localdate()
        expired = StockBatch.objects.create(
            shop=self.shop, product=self.product, quantity_received=Decimal("7"),
            quantity_remaining=Decimal("7"), expiry_date=today - timedelta(days=1),
        )
        soon = StockBatch.objects.create(
            shop=self.shop, product=self.product, quantity_received=Decimal("3"),
            quantity_remaining=Decimal("3"), expiry_date=today + timedelta(days=2),
        )
        later = StockBatch.objects.create(
            shop=self.shop, product=self.product, quantity_received=Decimal("4"),
            quantity_remaining=Decimal("4"), expiry_date=today + timedelta(days=8),
        )
        invoice = create_invoice(
            user=self.user, shop=self.shop, customer=None,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "4"}]' % self.product.pk,
        )
        soon.refresh_from_db()
        later.refresh_from_db()
        expired.refresh_from_db()
        self.assertEqual(soon.quantity_remaining, Decimal("0"))
        self.assertEqual(later.quantity_remaining, Decimal("3"))
        self.assertEqual(expired.quantity_remaining, Decimal("7"))
        self.assertEqual(list(InvoiceItemBatch.objects.filter(invoice_item__invoice=invoice).values_list("batch_id", flat=True)), [soon.pk, later.pk])

    def test_insufficient_stock_does_not_create_invoice_or_change_stock(self):
        with self.assertRaises(ValidationError):
            create_invoice(
                user=self.user, shop=self.shop, customer=None,
                payment_type=Invoice.PaymentType.CASH, amount_paid=0,
                items_json='[{"product_id": %s, "quantity": "11"}]' % self.product.pk,
            )
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.quantity_remaining, Decimal("10"))
        self.assertFalse(Invoice.objects.filter(shop=self.shop).exists())

    def test_cancelling_invoice_restores_stock_and_reverses_credit(self):
        invoice = create_invoice(
            user=self.user, shop=self.shop, customer=self.customer,
            payment_type=Invoice.PaymentType.CREDIT, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "3"}]' % self.product.pk,
        )
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.quantity_remaining, Decimal("7"))
        self.assertEqual(calculate_credit_balance(self.customer), Decimal("30.00"))
        cancel_invoice(invoice=invoice, user=self.user, shop=self.shop)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.quantity_remaining, Decimal("10"))
        self.assertEqual(calculate_credit_balance(self.customer), Decimal("0.00"))
        invoice.refresh_from_db()
        self.assertEqual(invoice.status, Invoice.Status.CANCELLED)

    def test_cancelling_invoice_does_not_make_credit_balance_negative_after_payment(self):
        invoice = create_invoice(
            user=self.user, shop=self.shop, customer=self.customer,
            payment_type=Invoice.PaymentType.CREDIT, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "3"}]' % self.product.pk,
        )
        CreditEntry.objects.create(
            shop=self.shop, customer=self.customer,
            entry_type=CreditEntry.EntryType.RECEIVED, amount=Decimal("10.00"),
        )

        cancel_invoice(invoice=invoice, user=self.user, shop=self.shop)

        self.assertEqual(calculate_credit_balance(self.customer), Decimal("0.00"))
        reversal = CreditEntry.objects.get(
            shop=self.shop, customer=self.customer, invoice=invoice,
            entry_type=CreditEntry.EntryType.RECEIVED,
        )
        self.assertEqual(reversal.amount, Decimal("20.00"))

    def test_credit_balance_is_given_minus_received(self):
        CreditEntry.objects.create(shop=self.shop, customer=self.customer, entry_type=CreditEntry.EntryType.GIVEN, amount=Decimal("100"))
        CreditEntry.objects.create(shop=self.shop, customer=self.customer, entry_type=CreditEntry.EntryType.RECEIVED, amount=Decimal("35"))
        self.assertEqual(calculate_credit_balance(self.customer), Decimal("65"))

    def test_public_invoice_renders_without_login(self):
        invoice = create_invoice(
            user=self.user, shop=self.shop, customer=self.customer,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "1"}]' % self.product.pk,
        )
        response = self.client.get(reverse("public_invoice", args=[invoice.share_token]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, invoice.number)
        self.assertContains(response, "Tax Invoice")
        self.assertContains(response, "Ten Rupees only")

    def test_about_and_developer_pages_require_login_and_show_bilingual_content(self):
        for route in ("about", "developer"):
            with self.subTest(route=route):
                response = self.client.get(reverse(route))
                self.assertRedirects(
                    response,
                    f"{reverse('login')}?next={reverse(route)}",
                )

        self.client.force_login(self.user)
        about_response = self.client.get(reverse("about"))
        self.assertEqual(about_response.status_code, 200)
        self.assertContains(about_response, "About ShopSetu")
        self.assertContains(about_response, "ShopSetu के बारे में")
        self.assertContains(about_response, "expiry dates")

        developer_response = self.client.get(reverse("developer"))
        self.assertEqual(developer_response.status_code, 200)
        self.assertContains(developer_response, "Rahul Kumar Prasad")
        self.assertContains(developer_response, "राहुल कुमार प्रसाद")
        self.assertContains(developer_response, "tel:+916203027876")
        self.assertContains(developer_response, "tel:+919508083649")
        self.assertContains(developer_response, "tel:+917061058012")
        self.assertContains(developer_response, "near City Hospital, Ranchi, Jharkhand 834005")

    def test_invoice_pdf_download_works_for_shop_and_public_share_link(self):
        invoice = create_invoice(
            user=self.user, shop=self.shop, customer=self.customer,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "1"}]' % self.product.pk,
        )
        self.client.force_login(self.user)
        for url in (
            reverse("invoice_pdf", args=[invoice.pk]),
            reverse("public_invoice_pdf", args=[invoice.share_token]),
        ):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Content-Type"], "application/pdf")
            self.assertIn(b"%PDF-", response.content[:10])

    def test_invoice_detail_offers_pdf_file_sharing_and_whatsapp_link_fallback(self):
        invoice = create_invoice(
            user=self.user, shop=self.shop, customer=self.customer,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "1"}]' % self.product.pk,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("invoice_detail", args=[invoice.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "data-invoice-share")
        self.assertContains(response, reverse("invoice_pdf", args=[invoice.pk]))
        self.assertContains(response, reverse("invoice_whatsapp", args=[invoice.pk]))
        self.assertContains(response, "Open WhatsApp with invoice link")

    def test_product_search_returns_matching_in_stock_shop_products(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("product_search"), {"q": "Tea"}, HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Tea")
        self.assertContains(response, "Add Tea to bill")

class InventoryQueryTests(ShopTestCase):
    def test_bulk_stock_import_previews_then_adds_batches_and_new_products(self):
        self.client.force_login(self.user)
        csv_content = (
            "product_name,unit,selling_price,purchase_price,quantity_received,expiry_date,gst_percent,hsn_code,low_stock_limit,note\n"
            "Tea,,,,5,,0,,,Restock\n"
            "Coffee,packet,25.00,18.00,12,2027-12-31,5,0901,3,Opening stock\n"
        )
        upload = SimpleUploadedFile("stock.csv", csv_content.encode("utf-8"), content_type="text/csv")

        preview_response = self.client.post(
            reverse("stock_bulk_import"),
            {"action": "preview", "spreadsheet": upload},
        )

        self.assertEqual(preview_response.status_code, 200)
        self.assertContains(preview_response, "Confirm and import all rows")
        self.assertContains(preview_response, "bulk-stock-page")
        self.assertContains(preview_response, "bulk-preview-table")
        self.assertEqual(Product.objects.filter(shop=self.shop).count(), 1)
        self.assertEqual(StockBatch.objects.filter(shop=self.shop).count(), 1)
        token = preview_response.context["preview_token"]
        confirm_response = self.client.post(
            reverse("stock_bulk_import"),
            {"action": "confirm", "preview_token": token},
        )

        self.assertRedirects(confirm_response, reverse("stock"))
        self.product.refresh_from_db()
        self.assertEqual(self.product.selling_price, Decimal("10.00"))
        tea_batch = StockBatch.objects.get(shop=self.shop, product=self.product, note="Restock")
        self.assertEqual(tea_batch.quantity_remaining, Decimal("5"))
        coffee = Product.objects.get(shop=self.shop, name="Coffee")
        coffee_batch = StockBatch.objects.get(shop=self.shop, product=coffee)
        self.assertEqual(coffee.unit, Product.Unit.PACKET)
        self.assertEqual(coffee_batch.quantity_remaining, Decimal("12"))
        self.assertEqual(coffee_batch.expiry_date.isoformat(), "2027-12-31")
        self.assertEqual(coffee_batch.purchase_price, Decimal("18.00"))

    def test_bulk_stock_import_blocks_invalid_rows_without_partial_import(self):
        self.client.force_login(self.user)
        csv_content = (
            "product_name,unit,selling_price,quantity_received\n"
            "Coffee,packet,,4\n"
            "Tea,,,2\n"
        )
        upload = SimpleUploadedFile("invalid.csv", csv_content.encode("utf-8"), content_type="text/csv")

        response = self.client.post(
            reverse("stock_bulk_import"),
            {"action": "preview", "spreadsheet": upload},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Selling price is required for a new product.")
        self.assertNotContains(response, "Confirm and import all rows")
        self.assertNotIn("preview_token", response.context)
        self.assertEqual(Product.objects.filter(shop=self.shop).count(), 1)
        self.assertEqual(StockBatch.objects.filter(shop=self.shop).count(), 1)

    def test_bulk_stock_import_preview_cannot_be_confirmed_for_another_shop(self):
        self.client.force_login(self.user)
        csv_content = (
            "product_name,unit,selling_price,quantity_received\n"
            "Coffee,packet,25.00,4\n"
        )
        upload = SimpleUploadedFile("stock.csv", csv_content.encode("utf-8"), content_type="text/csv")
        preview = self.client.post(
            reverse("stock_bulk_import"),
            {"action": "preview", "spreadsheet": upload},
        )
        other_owner = User.objects.create_user("other-owner")
        UserProfile.objects.create(
            user=other_owner, shop=self.other_shop, role=UserProfile.Role.OWNER,
            must_change_password=False,
        )
        self.client.force_login(other_owner)

        response = self.client.post(
            reverse("stock_bulk_import"),
            {"action": "confirm", "preview_token": preview.context["preview_token"]},
        )

        self.assertRedirects(response, reverse("stock_bulk_import"))
        self.assertFalse(Product.objects.filter(shop=self.shop, name="Coffee").exists())
        self.assertFalse(Product.objects.filter(shop=self.other_shop, name="Coffee").exists())

    def test_staff_cannot_open_bulk_stock_entry(self):
        staff = User.objects.create_user("bulk-staff")
        UserProfile.objects.create(
            user=staff, shop=self.shop, role=UserProfile.Role.STAFF,
            must_change_password=False,
        )
        self.client.force_login(staff)

        response = self.client.get(reverse("stock_bulk_import"))

        self.assertEqual(response.status_code, 403)

    def test_bulk_stock_import_reads_excel_workbooks(self):
        self.client.force_login(self.user)
        workbook = load_workbook(BytesIO(self.client.get(reverse("stock_bulk_template")).content))
        sheet = workbook["Stock import"]
        sheet.append(["Tea", "", "", "4.25", "3", "2027-12-31", "", "", "", "Excel restock"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        upload = SimpleUploadedFile(
            "restock.xlsx", output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            reverse("stock_bulk_import"),
            {"action": "preview", "spreadsheet": upload},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Confirm and import all rows")
        confirm_response = self.client.post(
            reverse("stock_bulk_import"),
            {"action": "confirm", "preview_token": response.context["preview_token"]},
        )
        self.assertRedirects(confirm_response, reverse("stock"))
        batch = StockBatch.objects.get(shop=self.shop, product=self.product, note="Excel restock")
        self.assertEqual(batch.quantity_remaining, Decimal("3"))
        self.assertEqual(batch.purchase_price, Decimal("4.25"))
        self.assertEqual(batch.expiry_date.isoformat(), "2027-12-31")

    def test_bulk_stock_template_is_an_excel_workbook_with_documented_headers(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("stock_bulk_template"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(
            tuple(cell.value for cell in workbook["Stock import"][1]),
            (
                "product_name", "unit", "selling_price", "purchase_price",
                "quantity_received", "expiry_date", "gst_percent", "hsn_code",
                "low_stock_limit", "note",
            ),
        )
        self.assertIn("Instructions", workbook.sheetnames)
        workbook.close()

    def test_dashboard_marks_low_stock_and_counts_expiry_windows(self):
        today = timezone.localdate()
        self.batch.quantity_remaining = Decimal("1")
        self.batch.save()
        StockBatch.objects.create(
            shop=self.shop, product=self.product, quantity_received=Decimal("1"),
            quantity_remaining=Decimal("1"), expiry_date=today + timedelta(days=5),
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("home"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "dashboard-metrics")
        self.assertContains(response, "mobile-dock")
        self.assertContains(response, "Developed &amp; engineered by Rahul Kumar Prasad")
        self.assertContains(response, "data-theme-toggle")
        self.assertContains(response, "metric-profit-today")
        self.assertContains(response, "Low stock")
        self.assertEqual(len(response.context["low_stock"]), 1)
        self.assertEqual(response.context["expiry_counts"][7], 1)
        self.assertEqual(response.context["expiry_counts"][15], 1)
        self.assertEqual(response.context["expiry_counts"][30], 1)

    def test_dashboard_profit_uses_known_batch_costs_and_excludes_uncosted_or_cancelled_bills(self):
        self.batch.purchase_price = Decimal("4.00")
        self.batch.save(update_fields=["purchase_price"])
        create_invoice(
            user=self.user, shop=self.shop, customer=None,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "2"}]' % self.product.pk,
        )
        uncosted_product = Product.objects.create(
            shop=self.shop, name="Unknown cost item", selling_price=Decimal("10.00"),
        )
        StockBatch.objects.create(
            shop=self.shop, product=uncosted_product, quantity_received=Decimal("5"),
            quantity_remaining=Decimal("5"),
        )
        create_invoice(
            user=self.user, shop=self.shop, customer=None,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "1"}]' % uncosted_product.pk,
        )
        cancelled_invoice = create_invoice(
            user=self.user, shop=self.shop, customer=None,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "1"}]' % self.product.pk,
        )
        cancel_invoice(invoice=cancelled_invoice, user=self.user, shop=self.shop)
        self.client.force_login(self.user)

        response = self.client.get(reverse("home"))

        self.assertEqual(response.context["profit_today"], Decimal("12.00"))
        self.assertEqual(response.context["sales_month"], Decimal("30.00"))
        self.assertEqual(response.context["profit_month"], Decimal("12.00"))
        self.assertContains(response, "bills without batch cost are excluded")
        self.assertContains(response, "Monthly sales")
        self.assertContains(response, "Monthly profit")

    def test_billing_recommends_all_time_best_sellers_from_active_bills(self):
        coffee = Product.objects.create(
            shop=self.shop, name="Coffee", selling_price=Decimal("20.00"),
        )
        StockBatch.objects.create(
            shop=self.shop, product=coffee, quantity_received=Decimal("10"),
            quantity_remaining=Decimal("10"),
        )
        tea_invoice = create_invoice(
            user=self.user, shop=self.shop, customer=None,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "3"}]' % self.product.pk,
        )
        create_invoice(
            user=self.user, shop=self.shop, customer=None,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "2"}]' % self.product.pk,
        )
        create_invoice(
            user=self.user, shop=self.shop, customer=None,
            payment_type=Invoice.PaymentType.CASH, amount_paid=0,
            items_json='[{"product_id": %s, "quantity": "4"}]' % coffee.pk,
        )
        cancel_invoice(invoice=tea_invoice, user=self.user, shop=self.shop)
        self.client.force_login(self.user)

        response = self.client.get(reverse("billing"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            [product.name for product in response.context["best_sellers"]],
            ["Coffee", "Tea"],
        )
        self.assertContains(response, "Most-selling products")
        self.assertContains(response, "Add Coffee to bill")
        self.assertContains(response, "data-product-quantity-controls")
        self.assertContains(response, 'data-quantity-step="-1"')
        self.assertContains(response, "bill-estimated-total")
        rendered = response.content.decode()
        self.assertLess(rendered.index('id="product-search"'), rendered.index("Bill summary"))
        self.assertLess(rendered.index("Bill summary"), rendered.index("Most-selling products"))

    def test_expired_batches_are_not_counted_as_upcoming(self):
        today = timezone.localdate()
        StockBatch.objects.create(
            shop=self.shop, product=self.product, quantity_received=Decimal("1"),
            quantity_remaining=Decimal("1"), expiry_date=today - timedelta(days=1),
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("home"))
        self.assertEqual(response.context["expiry_counts"][7], 0)
        self.assertEqual(response.context["expired"].count(), 1)

    def test_stock_correction_records_signed_quantity_change(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("stock_correct", args=[self.batch.pk]),
            {"direction": "INCREASE", "quantity_change": "2", "reason": "CORRECTION"},
        )
        self.assertEqual(response.status_code, 302)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.quantity_remaining, Decimal("12"))
        adjustment = self.batch.adjustments.get()
        self.assertEqual(adjustment.quantity_change, Decimal("2"))
        self.assertEqual(adjustment.shop_id, self.shop.pk)


class DemoCommandTests(TestCase):
    def test_create_demo_shop_builds_owner_and_sample_inventory(self):
        call_command("create_demo_shop", username="demo-test-owner", password="Demo-password-123")
        user = User.objects.get(username="demo-test-owner")
        self.assertEqual(user.profile.role, UserProfile.Role.OWNER)
        self.assertTrue(Product.objects.filter(shop=user.profile.shop).exists())
        self.assertTrue(StockBatch.objects.filter(shop=user.profile.shop).exists())
