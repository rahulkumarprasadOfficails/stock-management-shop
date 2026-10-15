import csv
import logging
import secrets
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth import logout as auth_logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import PasswordChangeForm
from django.contrib.auth.models import User
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.db.models.deletion import ProtectedError
from django.db.models import Count, F, Prefetch, Q, Sum
from django.http import HttpResponse
from django.http import JsonResponse
from django.templatetags.static import static
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils.translation import gettext as _
from django.views.decorators.http import require_POST
from django.contrib.auth.views import LoginView
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .bulk_stock import HEADERS, MAX_IMPORT_BYTES, read_stock_spreadsheet
from .forms import (
    BillingForm, BulkStockRowForm, CorrectionForm, CreditForm, CustomerForm, LoginForm, ProductForm,
    ProductStockForm, PromotionForm, ShopSetupForm, StaffForm,
)
from .models import (
    CreditEntry, Customer, Invoice, InvoiceItem, InvoiceItemBatch,
    Product, PromotionLog, Shop, StockAdjustment, StockBatch, UserProfile,
)
from .pdf import build_invoice_pdf
from .services import calculate_credit_balance, cancel_invoice, create_invoice


logger = logging.getLogger(__name__)


def pwa_manifest(_request):
    from django.http import JsonResponse
    return JsonResponse({
        "name": "ShopSetu",
        "short_name": "ShopSetu",
        "start_url": "/",
        "display": "standalone",
        "background_color": "#f8fafc",
        "theme_color": "#0f766e",
        "icons": [
            {"src": static("store/icon-192.svg"), "sizes": "192x192", "type": "image/svg+xml", "purpose": "any"},
            {"src": static("store/icon-512.svg"), "sizes": "512x512", "type": "image/svg+xml", "purpose": "any maskable"},
        ],
    })


def service_worker(_request):
    response = HttpResponse(
        "self.addEventListener('fetch', event => event.respondWith(fetch(event.request)));",
        content_type="application/javascript",
    )
    response["Service-Worker-Allowed"] = "/"
    response["Cache-Control"] = "no-cache"
    return response


def current_shop(request):
    if request.user.is_superuser:
        raise PermissionDenied
    return request.user.profile.shop


def owner_only(view):
    def wrapped(request, *args, **kwargs):
        if request.user.profile.role != UserProfile.Role.OWNER:
            raise PermissionDenied
        return view(request, *args, **kwargs)
    return login_required(wrapped)


class ShopLoginView(LoginView):
    template_name = "registration/login.html"
    authentication_form = LoginForm

    def form_valid(self, form):
        response = super().form_valid(form)
        target = self.get_success_url()
        if not url_has_allowed_host_and_scheme(target, {self.request.get_host()}):
            return redirect("home")
        return response


@login_required
def logout_view(request):
    if request.method == "POST":
        auth_logout(request)
    return redirect("login")


@login_required
def password_change(request):
    profile = request.user.profile
    form = PasswordChangeForm(request.user, request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save()
        profile.must_change_password = False
        profile.save(update_fields=["must_change_password"])
        messages.success(request, _("Your password has been changed."))
        return redirect("setup" if profile.role == UserProfile.Role.OWNER and not profile.shop.setup_completed else "home")
    return render(request, "registration/password_change.html", {"form": form})


def account_inactive(request):
    return render(request, "registration/account_inactive.html")


@login_required
def about(request):
    return render(request, "store/about.html")


@login_required
def developer(request):
    return render(request, "store/developer.html")


@owner_only
def setup(request):
    shop = current_shop(request)
    form = ShopSetupForm(request.POST or None, instance=shop)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        obj.setup_completed = True
        obj.save()
        messages.success(request, _("Shop setup is complete."))
        return redirect("home")
    return render(request, "store/setup.html", {"form": form})


@login_required
def home(request):
    shop = current_shop(request)
    if request.user.profile.role == UserProfile.Role.STAFF:
        return redirect("billing")
    today = timezone.localdate()
    products = Product.objects.for_user(request.user).filter(is_active=True).annotate(
        stock=Sum("batches__quantity_remaining", filter=Q(batches__shop=shop))
    )
    low_stock = [p for p in products if (p.stock or Decimal("0")) <= p.low_stock_limit]
    batches = StockBatch.objects.for_user(request.user).filter(quantity_remaining__gt=0, expiry_date__isnull=False)
    expiry_summary = batches.aggregate(
        **{
            f"days_{days}": Count(
                "id",
                filter=Q(expiry_date__gte=today, expiry_date__lte=today + timedelta(days=days)),
            )
            for days in (7, 15, 30)
        }
    )
    expiry_counts = {days: expiry_summary[f"days_{days}"] for days in (7, 15, 30)}
    expired = batches.filter(expiry_date__lt=today).select_related("product").order_by("expiry_date")
    month_start = today.replace(day=1)
    period_invoices = Invoice.objects.for_user(request.user).filter(
        status=Invoice.Status.ACTIVE, date__date__gte=month_start, date__date__lte=today,
    ).prefetch_related(
        Prefetch(
            "items",
            queryset=InvoiceItem.objects.filter(shop=shop).prefetch_related(
                Prefetch(
                    "batch_allocations",
                    queryset=InvoiceItemBatch.objects.filter(shop=shop).select_related("batch"),
                )
            ),
        )
    )
    sales_today = Decimal("0")
    sales_month = Decimal("0")
    profit_today = Decimal("0")
    profit_month = Decimal("0")
    for invoice in period_invoices:
        sales_month += invoice.total
        cost_of_goods = Decimal("0")
        fully_costed = True
        for item in invoice.items.all():
            allocations = list(item.batch_allocations.all())
            allocated_quantity = sum(
                (allocation.quantity for allocation in allocations), Decimal("0"),
            )
            if not allocations or allocated_quantity != item.quantity:
                fully_costed = False
                break
            for allocation in allocations:
                if allocation.batch.purchase_price is None:
                    fully_costed = False
                    break
                cost_of_goods += allocation.quantity * allocation.batch.purchase_price
            if not fully_costed:
                break
        if fully_costed:
            invoice_profit = invoice.subtotal - cost_of_goods
            profit_month += invoice_profit
        else:
            invoice_profit = None
        if timezone.localtime(invoice.date).date() == today:
            sales_today += invoice.total
            if invoice_profit is not None:
                profit_today += invoice_profit
    profit_today = profit_today.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    profit_month = profit_month.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    credit_summary = CreditEntry.objects.for_user(request.user).aggregate(
        given=Sum("amount", filter=Q(entry_type=CreditEntry.EntryType.GIVEN)),
        received=Sum("amount", filter=Q(entry_type=CreditEntry.EntryType.RECEIVED)),
    )
    pending_credit = (credit_summary["given"] or Decimal("0")) - (credit_summary["received"] or Decimal("0"))
    return render(request, "store/home.html", {
        "low_stock": low_stock, "expiry_counts": expiry_counts, "expired": expired,
        "sales_today": sales_today, "profit_today": profit_today,
        "sales_month": sales_month, "profit_month": profit_month,
        "pending_credit": pending_credit,
    })


@login_required
def product_list(request):
    current_shop(request)
    return redirect("stock")


@owner_only
def product_create(request):
    current_shop(request)
    return redirect("stock_add")


@owner_only
def product_edit(request, pk):
    shop = current_shop(request)
    product = get_object_or_404(Product, shop=shop, pk=pk)
    form = ProductForm(request.POST or None, instance=product, shop=shop)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, _("Product updated."))
        return redirect("products")
    return render(request, "store/product_form.html", {"form": form, "title": "Edit product"})


@login_required
def stock_list(request):
    shop = current_shop(request)
    query = request.GET.get("q", "").strip()
    products = Product.objects.for_user(request.user).filter(is_active=True).prefetch_related(
        Prefetch("batches", queryset=StockBatch.objects.for_user(request.user))
    )
    if query:
        products = products.filter(name__icontains=query)
    return render(request, "store/stock.html", {"products": products, "query": query})


@owner_only
@require_POST
def product_delete(request, pk):
    shop = current_shop(request)
    try:
        with transaction.atomic():
            product = get_object_or_404(
                Product.objects.for_user(request.user).select_for_update(),
                pk=pk,
                shop=shop,
            )
            has_invoice_history = InvoiceItem.objects.filter(
                shop=shop, product=product,
            ).exists()
            has_adjustment_history = StockAdjustment.objects.filter(
                shop=shop, batch__product=product,
            ).exists()
            if has_invoice_history or has_adjustment_history:
                messages.error(
                    request,
                    _("This product cannot be deleted because it has invoice or stock-adjustment history."),
                )
            else:
                product.delete()
                messages.success(request, _("Product and its stock batches were deleted."))
    except ProtectedError:
        messages.error(
            request,
            _("This product cannot be deleted because it is linked to existing records."),
        )
    return redirect("stock")


@owner_only
def stock_add(request):
    shop = current_shop(request)
    form = ProductStockForm(request.POST or None, shop=shop)
    if request.method == "POST" and form.is_valid():
        data = form.cleaned_data
        with transaction.atomic():
            product = Product.objects.create(
                shop=shop, name=data["name"], unit=data["unit"],
                selling_price=data["selling_price"], purchase_price=data["purchase_price"],
                gst_percent=int(data["gst_percent"]), hsn_code=data["hsn_code"],
                low_stock_limit=data["low_stock_limit"],
            )
            StockBatch.objects.create(
                shop=shop, product=product,
                quantity_received=data["quantity_received"],
                quantity_remaining=data["quantity_received"],
                expiry_date=data["expiry_date"], purchase_price=data["purchase_price"],
                note=data["note"],
            )
        messages.success(request, _("Product and opening stock added."))
        return redirect("stock")
    return render(request, "store/stock_add.html", {
        "form": form,
    })


@owner_only
def stock_bulk_template(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stock import"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for column, width in zip("ABCDEFGHIJ", (26, 14, 16, 16, 19, 16, 14, 16, 18, 28)):
        sheet.column_dimensions[column].width = width

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["ShopSetu bulk stock import"])
    instructions.append(["Required columns", "product_name, quantity_received"])
    instructions.append(["New product", "Also provide unit and selling_price. GST defaults to 0%."])
    instructions.append(["Existing product", "Match by product name; only a new stock batch is added. Product details are not changed."])
    instructions.append(["Optional batch fields", "purchase_price, expiry_date (YYYY-MM-DD), note"])
    instructions.append(["Optional new-product fields", "gst_percent, hsn_code, low_stock_limit"])
    instructions.append(["Units", ", ".join(value for value, _label in Product.Unit.choices)])
    instructions.append(["Limits", "Maximum 5,000 product rows and 5 MB per upload."])
    instructions.column_dimensions["A"].width = 26
    instructions.column_dimensions["B"].width = 100
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="shopsetu-stock-template.xlsx"'
    return response


@owner_only
def stock_bulk_import(request):
    shop = current_shop(request)
    context = {}
    if request.method == "POST" and request.POST.get("action") == "confirm":
        token = request.POST.get("preview_token", "")
        session_key = f"bulk_stock_preview_{token}"
        preview = request.session.get(session_key)
        if (
            not token
            or not preview
            or preview["user_id"] != request.user.pk
            or preview["shop_id"] != shop.pk
            or timezone.now().timestamp() - preview["created_at"] > 3600
        ):
            messages.error(request, _("This stock preview has expired. Upload the spreadsheet again."))
            return redirect("stock_bulk_import")
        try:
            with transaction.atomic():
                Shop.objects.select_for_update().get(pk=shop.pk)
                for row in preview["rows"]:
                    if row["product_id"] is None:
                        if Product.objects.filter(shop=shop, name__iexact=row["product_name"]).exists():
                            raise ValidationError(
                                _("Product '%(name)s' was added after this preview. Upload the spreadsheet again.") % {
                                    "name": row["product_name"],
                                }
                            )
                        product = Product.objects.create(
                            shop=shop,
                            name=row["product_name"],
                            unit=row["unit"],
                            selling_price=Decimal(row["selling_price"]),
                            purchase_price=(
                                Decimal(row["purchase_price"])
                                if row["purchase_price"] is not None else None
                            ),
                            gst_percent=row["gst_percent"],
                            hsn_code=row["hsn_code"],
                            low_stock_limit=Decimal(row["low_stock_limit"]),
                        )
                    else:
                        product = Product.objects.select_for_update().get(
                            shop=shop, pk=row["product_id"], is_active=True,
                        )
                    quantity = Decimal(row["quantity_received"])
                    StockBatch.objects.create(
                        shop=shop,
                        product=product,
                        quantity_received=quantity,
                        quantity_remaining=quantity,
                        expiry_date=date.fromisoformat(row["expiry_date"]) if row["expiry_date"] else None,
                        purchase_price=(
                            Decimal(row["purchase_price"])
                            if row["purchase_price"] is not None else None
                        ),
                        note=row["note"],
                    )
        except Product.DoesNotExist:
            messages.error(request, _("A product in this import is no longer active. Upload the spreadsheet again."))
            return redirect("stock_bulk_import")
        except IntegrityError:
            logger.exception("Bulk stock import conflicted while saving shop %s.", shop.pk)
            messages.error(request, _("The import could not be saved because product data changed. Upload the spreadsheet again."))
            return render(request, "store/stock_bulk_import.html", context, status=409)
        except ValidationError as exc:
            messages.error(request, " ".join(exc.messages))
            return redirect("stock_bulk_import")

        imported_count = len(preview["rows"])
        del request.session[session_key]
        messages.success(
            request,
            _("Imported %(count)s stock rows successfully. Existing product details were left unchanged.") % {
                "count": imported_count,
            },
        )
        return redirect("stock")

    if request.method == "POST":
        upload = request.FILES.get("spreadsheet")
        if upload is None:
            context["upload_error"] = _("Choose an .xlsx or .csv spreadsheet to upload.")
            return render(request, "store/stock_bulk_import.html", context, status=400)
        if upload.size > MAX_IMPORT_BYTES:
            context["upload_error"] = _("The spreadsheet must be 5 MB or smaller.")
            return render(request, "store/stock_bulk_import.html", context, status=400)
        try:
            source_rows = read_stock_spreadsheet(upload)
        except ValidationError as exc:
            context["upload_error"] = " ".join(exc.messages)
            return render(request, "store/stock_bulk_import.html", context, status=400)

        existing_products = {
            product.name.casefold(): product
            for product in Product.objects.filter(shop=shop)
        }
        seen_names = set()
        preview_rows = []
        prepared_rows = []
        for source_row in source_rows:
            form = BulkStockRowForm(
                data=source_row,
                existing_products=existing_products,
            )
            valid = form.is_valid()
            product_name = (source_row.get("product_name") or "").strip()
            errors = [
                f"{field}: {error}"
                for field, error_list in form.errors.items()
                for error in error_list
            ]
            normalized_name = product_name.casefold()
            if normalized_name and normalized_name in seen_names:
                errors.append(_("Product appears more than once in the spreadsheet."))
            if normalized_name:
                seen_names.add(normalized_name)

            row = {
                "row_number": source_row["row_number"],
                "product_name": product_name,
                "quantity_received": source_row.get("quantity_received", ""),
                "errors": errors,
            }
            if valid and not errors:
                row["product_id"] = form.cleaned_data["existing_product"].pk if form.cleaned_data["existing_product"] else None
            preview_rows.append(row)
            if valid and not errors:
                cleaned = form.cleaned_data
                existing = cleaned["existing_product"]
                prepared_rows.append({
                    "product_id": existing.pk if existing else None,
                    "product_name": cleaned["product_name"],
                    "unit": cleaned["unit"],
                    "selling_price": str(cleaned["selling_price"]) if not existing else None,
                    "purchase_price": (
                        str(cleaned["purchase_price"])
                        if cleaned["purchase_price"] is not None else None
                    ),
                    "quantity_received": str(cleaned["quantity_received"]),
                    "expiry_date": cleaned["expiry_date"].isoformat() if cleaned["expiry_date"] else "",
                    "gst_percent": cleaned["gst_percent"],
                    "hsn_code": cleaned["hsn_code"],
                    "low_stock_limit": str(cleaned["low_stock_limit"]),
                    "note": cleaned["note"],
                })
        if any(row["errors"] for row in preview_rows):
            context["preview_rows"] = preview_rows
            context["preview_error_count"] = sum(bool(row["errors"]) for row in preview_rows)
            context["can_confirm"] = False
        else:
            token = secrets.token_urlsafe(24)
            request.session[f"bulk_stock_preview_{token}"] = {
                "user_id": request.user.pk,
                "shop_id": shop.pk,
                "created_at": timezone.now().timestamp(),
                "rows": prepared_rows,
            }
            context.update({
                "preview_rows": preview_rows,
                "preview_token": token,
                "can_confirm": True,
            })
    return render(request, "store/stock_bulk_import.html", context)


@owner_only
def stock_correct(request, pk):
    shop = current_shop(request)
    batch = get_object_or_404(StockBatch.objects.for_user(request.user).select_related("product"), pk=pk)
    form = CorrectionForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        amount = form.cleaned_data["quantity_change"]
        change = amount if form.cleaned_data["direction"] == "INCREASE" else -amount
        with transaction.atomic():
            locked = StockBatch.objects.for_user(request.user).select_for_update().get(pk=batch.pk)
            if locked.quantity_remaining + change < 0:
                form.add_error("quantity_change", _("Cannot reduce stock below zero."))
            else:
                locked.quantity_remaining = F("quantity_remaining") + change
                locked.save(update_fields=["quantity_remaining"])
                StockAdjustment.objects.create(
                    shop=shop, batch=locked, quantity_change=change,
                    reason=form.cleaned_data["reason"], created_by=request.user,
                )
                messages.success(request, _("Stock correction recorded."))
                return redirect("stock")
    return render(request, "store/stock_correct.html", {"batch": batch, "form": form})


@login_required
def product_search(request):
    query = request.GET.get("q", "").strip()
    products = _billable_products(request.user, query=query)[:24]
    return render(request, "store/partials/product_search_results.html", {
        "products": products,
        "search_started": bool(query),
    })


def _billable_products(user, query=""):
    today = timezone.localdate()
    sellable_batches = Q(batches__shop_id=user.profile.shop_id, batches__quantity_remaining__gt=0) & (
        Q(batches__expiry_date__isnull=True) | Q(batches__expiry_date__gte=today)
    )
    products = Product.objects.for_user(user).filter(is_active=True).annotate(
        sellable_stock=Sum("batches__quantity_remaining", filter=sellable_batches)
    ).order_by("name")
    if query:
        products = products.filter(name__icontains=query)
    return products


@login_required
def billing(request):
    shop = current_shop(request)
    form = BillingForm(request.POST or None, shop=shop)
    if request.method == "POST" and form.is_valid():
        try:
            with transaction.atomic():
                customer = form.cleaned_data["customer"]
                if form.cleaned_data["new_customer_name"]:
                    customer = Customer.objects.create(
                        shop=shop,
                        name=form.cleaned_data["new_customer_name"],
                        whatsapp_number=form.cleaned_data["new_customer_whatsapp_number"],
                        address=form.cleaned_data["new_customer_address"],
                    )
                invoice = create_invoice(
                    user=request.user, shop=shop, customer=customer,
                    payment_type=form.cleaned_data["payment_type"],
                    amount_paid=form.cleaned_data.get("amount_paid"),
                    items_json=form.cleaned_data["items_json"],
                )
        except ValidationError as exc:
            form.add_error(None, exc.messages)
        else:
            return redirect("invoice_detail", pk=invoice.pk)
    return render(request, "store/billing.html", {
        "form": form,
        "today": timezone.localdate(),
        "best_sellers": _best_selling_products(request.user),
    })


def _best_selling_products(user, limit=5):
    product_sales = InvoiceItem.objects.for_user(user).filter(
        invoice__status=Invoice.Status.ACTIVE,
        product__is_active=True,
    ).values("product_id").annotate(
        sold_quantity=Sum("quantity"),
    ).order_by("-sold_quantity", "product_id")[:limit]
    ranked_ids = [row["product_id"] for row in product_sales]
    products = _billable_products(user).filter(pk__in=ranked_ids)
    products_by_id = {product.pk: product for product in products}
    return [products_by_id[pk] for pk in ranked_ids if pk in products_by_id]


@login_required
def invoice_detail(request, pk):
    shop = current_shop(request)
    invoice = get_object_or_404(
        Invoice.objects.for_user(request.user).select_related("shop", "customer", "created_by"),
        pk=pk,
    )
    invoice_items = InvoiceItem.objects.for_user(request.user).filter(
        invoice=invoice, product__shop=shop,
    ).select_related("product").prefetch_related("batch_allocations")
    share_url = request.build_absolute_uri(reverse("public_invoice", kwargs={"token": invoice.share_token}))
    return render(request, "store/invoice.html", {
        "invoice": invoice, "invoice_items": invoice_items, "share_url": share_url, "public": False,
    })


@login_required
def invoice_pdf(request, pk):
    shop = current_shop(request)
    invoice = get_object_or_404(Invoice.objects.for_user(request.user), pk=pk)
    items = InvoiceItem.objects.for_user(request.user).filter(
        invoice=invoice, product__shop=shop,
    ).select_related("product").order_by("id")
    return build_invoice_pdf(invoice, items)


def public_invoice(request, token):
    invoice = get_object_or_404(Invoice.objects.select_related("shop", "customer"), share_token=token)
    invoice_items = InvoiceItem.objects.filter(
        shop=invoice.shop, invoice=invoice, product__shop=invoice.shop,
    ).select_related("product")
    return render(request, "store/invoice.html", {
        "invoice": invoice, "invoice_items": invoice_items, "public": True, "share_url": "",
    })


def public_invoice_pdf(request, token):
    invoice = get_object_or_404(Invoice.objects.select_related("shop", "customer"), share_token=token)
    items = InvoiceItem.objects.filter(
        shop=invoice.shop, invoice=invoice, product__shop=invoice.shop,
    ).select_related("product").order_by("id")
    return build_invoice_pdf(invoice, items)


@owner_only
@require_POST
def invoice_cancel(request, pk):
    shop = current_shop(request)
    invoice = get_object_or_404(Invoice, shop=shop, pk=pk)
    try:
        cancel_invoice(invoice=invoice, user=request.user, shop=shop)
    except ValidationError as exc:
        messages.error(request, exc.messages[0])
    else:
        messages.success(request, "Invoice cancelled and stock restored.")
    return redirect("invoice_detail", pk=pk)


@login_required
def customer_list(request):
    shop = current_shop(request)
    query = request.GET.get("q", "").strip()
    customers = Customer.objects.for_user(request.user)
    if query:
        customers = customers.filter(Q(name__icontains=query) | Q(whatsapp_number__icontains=query))
    return render(request, "store/customers.html", {"customers": customers, "query": query})


@login_required
def customer_create(request):
    shop = current_shop(request)
    form = CustomerForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        customer = form.save(commit=False)
        customer.shop = shop
        customer.save()
        messages.success(request, _("Customer saved."))
        return redirect("customer_detail", pk=customer.pk)
    return render(request, "store/customer_form.html", {"form": form, "title": "Add customer"})


@login_required
def customer_edit(request, pk):
    shop = current_shop(request)
    customer = get_object_or_404(Customer, shop=shop, pk=pk)
    form = CustomerForm(request.POST or None, instance=customer)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, _("Customer updated."))
        return redirect("customer_detail", pk=pk)
    return render(request, "store/customer_form.html", {"form": form, "title": "Edit customer"})


@login_required
def customer_detail(request, pk):
    shop = current_shop(request)
    customer = get_object_or_404(Customer.objects.for_user(request.user), pk=pk)
    invoices = Invoice.objects.for_user(request.user).filter(customer=customer).order_by("-date")
    entries = CreditEntry.objects.for_user(request.user).filter(customer=customer)
    return render(request, "store/customer_detail.html", {
        "customer": customer, "invoices": invoices, "entries": entries,
        "balance": calculate_credit_balance(customer),
    })


@login_required
def credit_list(request):
    shop = current_shop(request)
    rows = []
    for customer in Customer.objects.for_user(request.user):
        balance = calculate_credit_balance(customer)
        if balance > 0:
            rows.append((customer, balance))
    return render(request, "store/credit_list.html", {"rows": rows})


@login_required
def credit_add(request, pk):
    shop = current_shop(request)
    customer = get_object_or_404(Customer, shop=shop, pk=pk)
    form = CreditForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            locked_customer = get_object_or_404(
                Customer.objects.for_user(request.user).select_for_update(),
                pk=pk,
                shop=shop,
            )
            entry_type = form.cleaned_data["entry_type"]
            amount = form.cleaned_data["amount"]
            balance = calculate_credit_balance(locked_customer)
            if entry_type == CreditEntry.EntryType.RECEIVED and amount > balance:
                form.add_error(
                    "amount",
                    _("Payment cannot exceed the pending credit of Rs. %(balance)s.")
                    % {"balance": f"{balance:.2f}"},
                )
            else:
                CreditEntry.objects.create(
                    shop=shop, customer=locked_customer, entry_type=entry_type,
                    amount=amount, note=form.cleaned_data["note"],
                )
                messages.success(request, _("Ledger entry saved."))
                return redirect("customer_detail", pk=pk)
    return render(request, "store/credit_form.html", {
        "form": form, "customer": customer,
        "balance": calculate_credit_balance(customer),
    })


@owner_only
def promotions(request):
    default_message = _("Hello {name}, we have fresh offers for you at our shop!")
    template = request.session.get("promotion_template", default_message)
    form = PromotionForm(request.POST or None, initial={"message": template})
    if request.method == "POST" and form.is_valid():
        request.session["promotion_template"] = form.cleaned_data["message"]
        return redirect("promotions")
    customers = Customer.objects.for_user(request.user).filter(is_promotion_allowed=True)
    return render(request, "store/promotions.html", {
        "form": form, "customers": customers, "promotion_message": template,
    })


@owner_only
@require_POST
def promotion_send(request, pk):
    shop = current_shop(request)
    customer = get_object_or_404(Customer.objects.for_user(request.user), pk=pk, is_promotion_allowed=True)
    template = request.session.get("promotion_template", _("Hello {name}, we have fresh offers for you at our shop!"))
    if request.POST.get("message"):
        form = PromotionForm(request.POST)
        if not form.is_valid():
            messages.error(request, _("The promotion message is invalid. Please save the template and try again."))
            return redirect("promotions")
        template = form.cleaned_data["message"]
    message = template.replace("{name}", customer.name)
    PromotionLog.objects.create(shop=shop, customer=customer, message=message)
    if request.headers.get("Accept") == "application/json":
        return JsonResponse({"status": "shared"})
    url = f"https://wa.me/{customer.whatsapp_international}?{urlencode({'text': message})}"
    return redirect(url)


@login_required
def invoice_whatsapp(request, pk):
    shop = current_shop(request)
    invoice = get_object_or_404(Invoice.objects.for_user(request.user).select_related("customer"), pk=pk)
    if not invoice.customer:
        messages.error(request, _("This invoice has no customer phone number."))
        return redirect("invoice_detail", pk=pk)
    link = request.build_absolute_uri(reverse("public_invoice", kwargs={"token": invoice.share_token}))
    message = f"Hello {invoice.customer.name}, here is your invoice {invoice.number}: {link}"
    return redirect(f"https://wa.me/{invoice.customer.whatsapp_international}?{urlencode({'text': message})}")


@login_required
def credit_whatsapp(request, pk):
    shop = current_shop(request)
    customer = get_object_or_404(Customer.objects.for_user(request.user), pk=pk)
    balance = calculate_credit_balance(customer)
    message = f"Hello {customer.name}, your pending balance at {shop.name} is Rs. {balance:.2f}."
    return redirect(f"https://wa.me/{customer.whatsapp_international}?{urlencode({'text': message})}")


@owner_only
def staff_list(request):
    shop = current_shop(request)
    staff = UserProfile.objects.for_user(request.user).filter(role=UserProfile.Role.STAFF).select_related("user")
    return render(request, "store/staff.html", {"staff": staff})


@owner_only
def staff_create(request):
    shop = current_shop(request)
    form = StaffForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            user = User.objects.create_user(
                username=form.cleaned_data["username"], password=form.cleaned_data["password"],
                is_active=True,
            )
            UserProfile.objects.create(user=user, shop=shop, role=UserProfile.Role.STAFF, must_change_password=True)
        messages.success(request, _("Staff login created. Share the temporary password securely."))
        return redirect("staff")
    return render(request, "store/staff_form.html", {"form": form})


@owner_only
@require_POST
def staff_toggle(request, pk):
    shop = current_shop(request)
    profile = get_object_or_404(UserProfile, shop=shop, role=UserProfile.Role.STAFF, pk=pk)
    profile.user.is_active = not profile.user.is_active
    profile.user.save(update_fields=["is_active"])
    messages.success(request, _("Staff login status updated."))
    return redirect("staff")


@owner_only
def reports(request):
    shop = current_shop(request)
    today = timezone.localdate()
    from django.utils.dateparse import parse_date
    start = parse_date(request.GET.get("start", "")) or today - timedelta(days=30)
    end = parse_date(request.GET.get("end", "")) or today
    invoices = Invoice.objects.for_user(request.user).filter(status=Invoice.Status.ACTIVE, date__date__gte=start, date__date__lte=end)
    sales = invoices.aggregate(total=Sum("total"))["total"] or Decimal("0")
    top_products = (
        InvoiceItem.objects.for_user(request.user).filter(invoice__in=invoices, invoice__shop=shop)
        .values("product__name").annotate(quantity=Sum("quantity"), total=Sum("line_total"))
        .order_by("-quantity")[:10]
    )
    expiring = StockBatch.objects.for_user(request.user).filter(
        quantity_remaining__gt=0, expiry_date__gte=today,
        expiry_date__lte=today + timedelta(days=30),
    ).select_related("product").order_by("expiry_date")
    dues = [(customer, calculate_credit_balance(customer)) for customer in Customer.objects.for_user(request.user)]
    dues = [(customer, amount) for customer, amount in dues if amount > 0]
    return render(request, "store/reports.html", {
        "sales": sales, "top_products": top_products, "expiring": expiring,
        "dues": dues, "start": start.isoformat(), "end": end.isoformat(),
    })


@owner_only
def report_csv(request, kind):
    shop = current_shop(request)
    today = timezone.localdate()
    from django.utils.dateparse import parse_date
    start = parse_date(request.GET.get("start", "")) or today - timedelta(days=30)
    end = parse_date(request.GET.get("end", "")) or today
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="shopsetu-{kind}.csv"'
    writer = csv.writer(response)
    today = timezone.localdate()
    if kind == "sales":
        writer.writerow(["Invoice", "Date", "Customer", "Total", "Payment", "Status"])
        invoices = Invoice.objects.for_user(request.user).filter(date__date__gte=start, date__date__lte=end)
        for row in invoices.iterator():
            writer.writerow([row.number, row.date.isoformat(), row.customer.name if row.customer else "Walk-in", row.total, row.payment_type, row.status])
    elif kind == "top-products":
        writer.writerow(["Product", "Quantity sold", "Sales"])
        rows = InvoiceItem.objects.for_user(request.user).filter(invoice__shop=shop, invoice__status=Invoice.Status.ACTIVE, invoice__date__date__gte=start, invoice__date__date__lte=end).values("product__name").annotate(quantity=Sum("quantity"), total=Sum("line_total")).order_by("-quantity")
        for row in rows.iterator():
            writer.writerow([row["product__name"], row["quantity"], row["total"]])
    elif kind == "expiring":
        writer.writerow(["Product", "Quantity", "Expiry date"])
        rows = StockBatch.objects.for_user(request.user).filter(quantity_remaining__gt=0, expiry_date__lte=today + timedelta(days=30)).select_related("product")
        for row in rows.iterator():
            writer.writerow([row.product.name, row.quantity_remaining, row.expiry_date])
    elif kind == "credit":
        writer.writerow(["Customer", "WhatsApp", "Pending"])
        for customer in Customer.objects.for_user(request.user).iterator():
            balance = calculate_credit_balance(customer)
            if balance:
                writer.writerow([customer.name, customer.whatsapp_number, balance])
    else:
        return HttpResponse("Unknown report.", status=404)
    return response
