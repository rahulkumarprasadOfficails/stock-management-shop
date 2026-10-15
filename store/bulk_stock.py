import csv
import io
import zipfile
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.utils.translation import gettext as _
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_IMPORT_ROWS = 5000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
HEADERS = (
    "product_name",
    "unit",
    "selling_price",
    "purchase_price",
    "quantity_received",
    "expiry_date",
    "gst_percent",
    "hsn_code",
    "low_stock_limit",
    "note",
)


def _normalize_header(value):
    return str(value or "").strip().lower().replace(" ", "_")


def _cell_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return "" if value is None else str(value).strip()


def read_stock_spreadsheet(upload):
    suffix = upload.name.rsplit(".", 1)[-1].lower() if "." in upload.name else ""
    try:
        if suffix == "csv":
            content = upload.read().decode("utf-8-sig")
            reader = csv.reader(io.StringIO(content, newline=""))
            headers_row = next(reader, [])
            sheet_rows = [headers_row]
            for values in reader:
                if any(value is not None and str(value).strip() for value in values):
                    if len(sheet_rows) > MAX_IMPORT_ROWS:
                        raise ValidationError(_("A single import can contain at most %(count)s product rows.") % {"count": MAX_IMPORT_ROWS})
                    sheet_rows.append(values)
        elif suffix == "xlsx":
            workbook = load_workbook(upload, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                if sheet.max_row > MAX_IMPORT_ROWS + 1:
                    raise ValidationError(_("A single import can contain at most %(count)s product rows.") % {"count": MAX_IMPORT_ROWS})
                sheet_rows = list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
        else:
            raise ValidationError(_("Upload an .xlsx or .csv file."))
    except (UnicodeDecodeError, csv.Error, InvalidFileException, zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ValidationError(_("The spreadsheet could not be read. Check that it is a valid .xlsx or UTF-8 CSV file.")) from exc

    while sheet_rows and not any(value is not None and str(value).strip() for value in sheet_rows[-1]):
        sheet_rows.pop()
    if not sheet_rows:
        raise ValidationError(_("The spreadsheet is empty."))

    headers = [_normalize_header(value) for value in sheet_rows[0]]
    nonempty_headers = [header for header in headers if header]
    if len(nonempty_headers) != len(set(nonempty_headers)):
        raise ValidationError(_("The spreadsheet has duplicate column headings."))
    missing_headers = {"product_name", "quantity_received"} - set(headers)
    if missing_headers:
        raise ValidationError(_("Required columns are missing: Product name and Quantity received."))

    rows = []
    for row_number, values in enumerate(sheet_rows[1:], start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ValidationError(_("A single import can contain at most %(count)s product rows.") % {"count": MAX_IMPORT_ROWS})
        if len(values) > len(headers):
            raise ValidationError(_("Row %(row)s has more cells than the header row.") % {"row": row_number})
        padded = list(values) + [None] * (len(headers) - len(values))
        rows.append({
            "row_number": row_number,
            **{header: _cell_value(value) for header, value in zip(headers, padded) if header},
        })
    if not rows:
        raise ValidationError(_("The spreadsheet has no product rows."))
    return rows
